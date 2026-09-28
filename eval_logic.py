"""
eval_logic.py
연말 인사 평가 계산 로직 (화면과 분리된 순수 계산 모듈)

핵심 규칙
1. 인원: 팀장 1명 + 팀원 7명 = 8명 고정
2. 원점수: 5개 영역 점수 합계 (100점 만점)
3. 팀장 원점수 = 팀원 7명 원점수의 중간값(Median)
   → 이후 8명 전체를 같은 비율로 환산하므로,
     환산 후에도 팀장 점수 = 팀원 7명 환산 점수의 중간값이 그대로 유지됨
4. 최종 환산 점수 = (본인 원점수 / 8명 원점수 합계) × 100
   → 소수 둘째 자리로 반올림하되, 합계가 정확히 100.00이 되도록 보정
"""

from io import BytesIO
import statistics

import pandas as pd

# ─────────────────────────────────────────────
# 1. 평가 기준 정의 (컬럼명, 만점, 설명)
# ─────────────────────────────────────────────
CRITERIA = [
    ("사업 규모", 25, "사업 담당 개수 및 규모 — 연간 담당 프로젝트 규모 및 난도"),
    ("완성도", 25, "사업 완성도 및 섬세함 — 지시 이행도, 마감 준수, 오류 최소화"),
    ("신규 발굴", 20, "신규 문의 응대 및 발굴 — 신규 사업 제안, 후원사 컨택, 문의 응대"),
    ("성실·근태", 15, "성실성/근태 — 지각·결근 및 기본 태도"),
    ("적극성", 15, "적극성 및 아이디어 — 능동적 태도, 아이디어 제안·실행"),
]
SCORE_COLS = [name for name, _, _ in CRITERIA]
MAX_SCORES = {name: mx for name, mx, _ in CRITERIA}

TEAM_SIZE = 7          # 팀원 수 (팀장 제외)
TOTAL_POINTS = 100.0   # 8명 최종 점수 합계
STAR = "⭐"

# 입력 표의 컬럼 순서 (엑셀 '입력데이터' 시트도 이 형식)
INPUT_COLS = ["이름", "직책", "주요 담당 사업"] + SCORE_COLS + ["팀장 총평 / 코멘트"]


def default_members() -> pd.DataFrame:
    """팀원 7명 입력용 빈 표(예시 이름 포함)를 만든다."""
    rows = []
    for i in range(1, TEAM_SIZE + 1):
        row = {"이름": f"팀원{i}", "직책": "", "주요 담당 사업": ""}
        row.update({c: 0.0 for c in SCORE_COLS})  # 0.5점 단위 입력을 위해 실수형
        row["팀장 총평 / 코멘트"] = ""
        rows.append(row)
    return pd.DataFrame(rows, columns=INPUT_COLS)


def clean_scores(df: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    """
    점수 컬럼을 숫자로 바꾸고 0~만점 범위로 제한한다.
    범위를 벗어난 값이 있으면 경고 메시지를 함께 돌려준다.
    """
    df = df.copy()
    warnings = []
    for col in SCORE_COLS:
        df[col] = pd.to_numeric(df[col], errors="coerce").fillna(0)
        mx = MAX_SCORES[col]
        bad = df[(df[col] < 0) | (df[col] > mx)]
        for _, r in bad.iterrows():
            warnings.append(f"{r['이름']}의 '{col}' 점수({r[col]:g})가 범위(0~{mx})를 벗어나 보정했습니다.")
        df[col] = df[col].clip(0, mx)
    for col in ["이름", "직책", "주요 담당 사업", "팀장 총평 / 코멘트"]:
        df[col] = df[col].fillna("").astype(str)
    return df, warnings


def round_to_exact_total(values: list[float], total: float = TOTAL_POINTS, decimals: int = 2) -> list[float]:
    """
    최대 잔여법(Largest Remainder Method) + 동점 유지:
    각 값을 소수 decimals 자리로 '내림'한 뒤, 모자란 0.01점들을
    잘려 나간 부분(잔여)이 큰 순서대로 나눠 주어 합계를 정확히 맞춘다.

    원래 값이 같은 사람들(예: 팀장과 중간값 팀원)은 한 묶음으로 처리해
    반올림 후에도 점수가 같게 유지한다. 묶음 전체에 줄 만큼 남지 않으면
    그 묶음은 건너뛰고 다음 사람에게 준다.
    """
    unit = 10 ** decimals
    target = round(total * unit)                     # 예: 100.00점 → 10000
    scaled = [v * unit for v in values]
    floored = [int(s + 1e-9) for s in scaled]        # 내림 (부동소수 오차 보정)
    shortfall = target - sum(floored)                # 부족한 0.01 단위 개수

    # 원래 값이 같은 사람끼리 묶기
    groups: dict[float, list[int]] = {}
    for i, s in enumerate(scaled):
        groups.setdefault(round(s, 6), []).append(i)
    # 잔여가 큰 묶음부터
    ordered = sorted(groups.values(), key=lambda g: scaled[g[0]] - floored[g[0]], reverse=True)

    for g in ordered:
        if shortfall >= len(g):
            for i in g:
                floored[i] += 1
            shortfall -= len(g)
    # 그래도 남으면(드문 경우) 동점이 아닌 사람에게 먼저, 이미 받은 사람은 뒤로 미뤄 개별 배분
    if shortfall > 0:
        group_size = {i: len(g) for g in groups.values() for i in g}
        singles = sorted(
            range(len(values)),
            key=lambda i: (group_size[i] > 1, floored[i] > int(scaled[i] + 1e-9), -(scaled[i] - int(scaled[i] + 1e-9))),
        )
        for i in singles[:shortfall]:
            floored[i] += 1
    return [f / unit for f in floored]


def compute_results(
    leader: dict,
    members: pd.DataFrame,
    top_n: int = 1,
    min_raw: float = 0.0,
    include_leader_in_promo: bool = False,
) -> pd.DataFrame:
    """
    평가 결과표를 계산한다.

    leader  : {"이름", "직책", "주요 담당 사업", "팀장 총평 / 코멘트"} (점수는 자동 산정)
    members : 팀원 7명 입력 표 (INPUT_COLS 형식)
    top_n   : 승진 추천 인원 (상위 N명)
    min_raw : 승진 추천 최소 원점수 (100점 만점 기준, 0이면 조건 없음)
    include_leader_in_promo : 팀장도 승진 추천 대상에 포함할지 여부
    """
    members, _ = clean_scores(members)
    members = members.copy()
    members["구분"] = "팀원"
    members["원점수 합계"] = members[SCORE_COLS].sum(axis=1)

    # ── 팀장 원점수 = 팀원 7명 원점수의 중간값 ──
    leader_raw = statistics.median(members["원점수 합계"].tolist())
    leader_row = {
        "이름": leader.get("이름", "팀장"),
        "직책": leader.get("직책", "팀장"),
        "주요 담당 사업": leader.get("주요 담당 사업", ""),
        "팀장 총평 / 코멘트": leader.get("팀장 총평 / 코멘트", ""),
        "구분": "팀장",
        "원점수 합계": leader_raw,
    }
    # 팀장은 영역별 점수가 없으므로 표시용으로 비워 둠
    leader_row.update({c: None for c in SCORE_COLS})

    df = pd.concat([pd.DataFrame([leader_row]), members], ignore_index=True)

    # ── 100점 총량 환산 ──
    raw_total = df["원점수 합계"].sum()
    if raw_total > 0:
        shares = (df["원점수 합계"] / raw_total * TOTAL_POINTS).tolist()
    else:
        # 모두 0점이면 균등 배분 (100 / 8 = 12.5)
        shares = [TOTAL_POINTS / len(df)] * len(df)
    df["최종 환산 점수"] = round_to_exact_total(shares)

    # ── 정렬 및 순위 (동점은 같은 순위) ──
    df = df.sort_values(["최종 환산 점수", "구분"], ascending=[False, False]).reset_index(drop=True)
    df["순위"] = df["최종 환산 점수"].rank(method="min", ascending=False).astype(int)

    # ── 승진 추천(⭐) 표시 ──
    candidates = df if include_leader_in_promo else df[df["구분"] == "팀원"]
    candidates = candidates[candidates["원점수 합계"] >= min_raw]
    if top_n > 0 and len(candidates) > 0:
        # 상위 N번째 점수와 동점인 사람까지 함께 추천
        cutoff = candidates["최종 환산 점수"].nlargest(min(top_n, len(candidates))).min()
        promo_idx = candidates[candidates["최종 환산 점수"] >= cutoff].index
    else:
        promo_idx = []
    df["승진 추천"] = ""
    df.loc[promo_idx, "승진 추천"] = STAR

    ordered = (
        ["승진 추천", "순위", "이름", "직책", "구분", "주요 담당 사업"]
        + SCORE_COLS
        + ["원점수 합계", "최종 환산 점수", "팀장 총평 / 코멘트"]
    )
    return df[ordered]


def to_excel_bytes(result: pd.DataFrame, leader: dict, members: pd.DataFrame) -> bytes:
    """
    결과를 서식이 적용된 엑셀 파일(bytes)로 만든다.
    시트 구성: ① 평가 결과  ② 입력데이터(다시 불러오기용)  ③ 평가 기준
    """
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    # 다시 불러오기용 입력 데이터 (첫 행 = 팀장, 점수는 비워 둠)
    leader_input = {c: leader.get(c, "") for c in INPUT_COLS}
    leader_input.update({c: None for c in SCORE_COLS})
    input_df = pd.concat([pd.DataFrame([leader_input]), members[INPUT_COLS]], ignore_index=True)
    input_df.insert(0, "구분", ["팀장"] + ["팀원"] * len(members))

    criteria_df = pd.DataFrame(
        [(n, m, d) for n, m, d in CRITERIA] + [("합계", 100, "")],
        columns=["영역", "배점", "설명"],
    )

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        result.to_excel(writer, sheet_name="평가 결과", index=False)
        input_df.to_excel(writer, sheet_name="입력데이터", index=False)
        criteria_df.to_excel(writer, sheet_name="평가 기준", index=False)

        header_fill = PatternFill("solid", fgColor="2F5D50")
        star_fill = PatternFill("solid", fgColor="FFF4CC")
        leader_fill = PatternFill("solid", fgColor="EEF3F1")
        thin = Side(style="thin", color="C8C8C8")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for ws in writer.sheets.values():
            # 헤더 서식
            for cell in ws[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            # 열 너비 자동 조정 (한글은 2칸으로 계산)
            for col_cells in ws.columns:
                letter = get_column_letter(col_cells[0].column)
                longest = max(
                    sum(2 if ord(ch) > 127 else 1 for ch in str(c.value or "")) for c in col_cells
                )
                ws.column_dimensions[letter].width = min(max(8, longest + 2), 50)
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    cell.border = border
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.freeze_panes = "A2"

        # 평가 결과 시트: 승진 추천 행 노란색, 팀장 행 연녹색
        ws = writer.sheets["평가 결과"]
        headers = [c.value for c in ws[1]]
        star_col = headers.index("승진 추천") + 1
        type_col = headers.index("구분") + 1
        score_col = headers.index("최종 환산 점수") + 1
        for r in range(2, ws.max_row + 1):
            fill = None
            if ws.cell(r, type_col).value == "팀장":
                fill = leader_fill
            if ws.cell(r, star_col).value == STAR:
                fill = star_fill
                ws.cell(r, star_col).alignment = Alignment(horizontal="center")
            if fill:
                for c in range(1, ws.max_column + 1):
                    ws.cell(r, c).fill = fill
            ws.cell(r, score_col).number_format = "0.00"
            ws.cell(r, score_col).font = Font(bold=True)
        # 합계 행 추가
        total_row = ws.max_row + 1
        ws.cell(total_row, score_col - 1, "합계").font = Font(bold=True)
        ws.cell(total_row, score_col, f"=SUM({get_column_letter(score_col)}2:{get_column_letter(score_col)}{total_row - 1})")
        ws.cell(total_row, score_col).number_format = "0.00"
        ws.cell(total_row, score_col).font = Font(bold=True)

    return buf.getvalue()


def load_from_excel(file) -> tuple[dict, pd.DataFrame]:
    """이전에 내려받은 엑셀의 '입력데이터' 시트를 읽어 팀장 정보와 팀원 표로 돌려준다."""
    df = pd.read_excel(file, sheet_name="입력데이터").fillna("")
    leader_rows = df[df["구분"] == "팀장"]
    member_rows = df[df["구분"] == "팀원"].head(TEAM_SIZE)
    leader = leader_rows.iloc[0][INPUT_COLS].to_dict() if len(leader_rows) else {}
    members = member_rows[INPUT_COLS].reset_index(drop=True)
    for c in SCORE_COLS:
        members[c] = pd.to_numeric(members[c], errors="coerce").fillna(0)
    return leader, members
